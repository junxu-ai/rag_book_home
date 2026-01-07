import json
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Dict, List, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


EMU_PER_INCH = 914400


# -------------------------------------------------
# Data model
# -------------------------------------------------
@dataclass
class Node:
    id: str
    name: str
    department: str
    title: str
    reports_to: Optional[str] = None
    children: List["Node"] = field(default_factory=list)
    depth: int = 0
    x_center: float = 0.0
    y_top: float = 0.0


# -------------------------------------------------
# Parsing
# -------------------------------------------------
def parse_flat_nodes(nodes_json: List[dict]) -> Node:
    nodes: Dict[str, Node] = {}
    for n in nodes_json:
        nodes[n["id"]] = Node(
            id=n["id"],
            name=n["name"],
            department=n["department"],
            title=n["title"],
            reports_to=n.get("reports_to"),
        )

    roots = []
    for n in nodes.values():
        if n.reports_to and n.reports_to in nodes:
            nodes[n.reports_to].children.append(n)
        else:
            roots.append(n)

    if len(roots) != 1:
        raise ValueError("Exactly one root node is required.")
    return roots[0]


def assign_depths(node: Node, depth: int = 0):
    node.depth = depth
    for c in node.children:
        assign_depths(c, depth + 1)


def collect_by_depth(root: Node) -> Dict[int, List[Node]]:
    levels = defaultdict(list)

    def dfs(n: Node):
        levels[n.depth].append(n)
        for c in n.children:
            dfs(c)

    dfs(root)
    return levels


# -------------------------------------------------
# Rendering helpers
# -------------------------------------------------
def add_three_row_box(slide, x, y, w, h, node: Node):
    outer = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    outer.fill.background()
    outer.line.color.rgb = RGBColor(120, 120, 120)

    row_h = h / 3.0

    rows = [
        (node.name, RGBColor(52, 73, 94), True),
        (node.department, RGBColor(93, 109, 126), False),
        (node.title, RGBColor(149, 165, 166), False),
    ]

    shapes = []

    for i, (text, color, bold) in enumerate(rows):
        r = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y + i * row_h),
            Inches(w),
            Inches(row_h),
        )
        r.fill.solid()
        r.fill.fore_color.rgb = color
        r.line.fill.background()

        tf = r.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.size = Pt(12 if i == 0 else 10)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(255, 255, 255)

        shapes.append(r)

    return outer


def add_connector(slide, parent, child):
    x1 = (parent.left + parent.width / 2) / EMU_PER_INCH
    y1 = (parent.top + parent.height) / EMU_PER_INCH
    x2 = (child.left + child.width / 2) / EMU_PER_INCH
    y2 = child.top / EMU_PER_INCH

    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = RGBColor(120, 120, 120)
    conn.line.width = Pt(1)


# -------------------------------------------------
# Main rendering
# -------------------------------------------------
def render_org_chart(json_path: str, output_path: str):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    root = parse_flat_nodes(data["nodes"])
    assign_depths(root)

    layer_alignment = data.get("layer_alignment", {})

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    box_w = 2.4
    box_h = 1.2
    level_gap = 1.6
    margin = 0.6

    levels = collect_by_depth(root)
    shape_map = {}

    for depth, nodes in levels.items():
        align = layer_alignment.get(str(depth), "center")
        count = len(nodes)
        total_w = count * box_w + (count - 1) * 0.4

        if align == "left":
            start_x = margin
        elif align == "right":
            start_x = 13.33 - margin - total_w
        else:  # center
            start_x = (13.33 - total_w) / 2

        y = margin + depth * level_gap

        for i, n in enumerate(nodes):
            x = start_x + i * (box_w + 0.4)
            shape = add_three_row_box(slide, x, y, box_w, box_h, n)
            shape_map[n.id] = shape

    for n in shape_map:
        node = next(x for x in data["nodes"] if x["id"] == n)
        if node.get("reports_to"):
            add_connector(slide, shape_map[node["reports_to"]], shape_map[n])

    prs.save(output_path)


if __name__ == "__main__":
    render_org_chart("pptx_json_flat.json", "org_chart.pptx")
